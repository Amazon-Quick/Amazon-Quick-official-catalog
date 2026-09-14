"""
description: Extracts and batches translatable text from DOCX and PPTX documents by walking the document structure and merging runs with identical formatting for parallel translation.
last_updated: 2026-09-13
origin: original

Phase 1, Document extraction for the translate-ms skill.

Offloads XML parsing, run merging, and state management to deterministic code
(rather than LLM reasoning) eliminates expensive inference steps, enables reliable
handling of large documents, and produces consistent results regardless of model context limits.

Provides extract_document() which detects .docx/.pptx, walks the document structure
(body, tables, headers/footers, shapes, notes), merges runs with identical formatting,
and batches paragraphs for parallel translation by workers. Also provides
get_first_words() for language detection.

Loaded in Phase 1 before spawning workers. Receives _STATE from state.py.

File layout:
  scripts/
    state.py         (TranslationState class + _STATE singleton)
    extract.py       (this file: extraction for both DOCX and PPTX)
    glossary.py      (custom glossary CSV parsing and matching)
    reconstruct.py   (worker result parsing + document reconstruction)
"""

from __future__ import annotations

import logging
import os
import sys
from dataclasses import dataclass
from enum import StrEnum
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from pptx import Presentation

# try/except avoids F821 lint errors in Quick's script concatenation mode.
try:
    from state import _STATE
except ImportError:
    pass  # Name provided by script concatenation at runtime


class _PrintStream:
    """Stdout adapter so logging records reach the real stdout via print()."""

    def write(self, message: str) -> int:
        if message:
            print(message, end="", file=sys.__stdout__)
        return len(message)

    def flush(self) -> None:
        return None


logging.basicConfig(
    stream=_PrintStream(),
    level=logging.INFO,
    format="%(asctime)s %(message)s",
    force=True,
)

_LOGGER = logging.getLogger(__name__)


WORDS_PER_BATCH = 250
MAX_PARAS_PER_BATCH = 30

_DOCX_EXTENSION = ".docx"
_PPTX_EXTENSION = ".pptx"
_DOCX_RUN_TAG = "r"
_DOCX_HYPERLINK_TAG = "hyperlink"
_PPTX_GROUP_SHAPE_TYPE = 6  # MSO_SHAPE_TYPE.GROUP
_DEFAULT_SAMPLE_WORDS = 100


class DocFormat(StrEnum):
    """Supported source document formats."""

    DOCX = "docx"
    PPTX = "pptx"


@dataclass(frozen=True)
class _MergedRun:
    """A run of text merged from one or more original runs with identical formatting."""

    text: str
    original_indices: list


def _get_all_runs_in_order(paragraph: Any) -> list:
    """Get ALL run elements including those inside <w:hyperlink> elements."""
    runs = []
    for child in paragraph._p:
        tag = etree.QName(child.tag).localname if isinstance(child.tag, str) else ""
        if tag == _DOCX_RUN_TAG:
            runs.append(child)
        elif tag == _DOCX_HYPERLINK_TAG:
            for sub in child:
                sub_tag = (
                    etree.QName(sub.tag).localname if isinstance(sub.tag, str) else ""
                )
                if sub_tag == _DOCX_RUN_TAG:
                    runs.append(sub)
    return runs


class _ParagraphBatcher:
    """Pure logic: group extracted paragraphs into batches by word count."""

    def __init__(self, words_per_batch: int, max_paras_per_batch: int) -> None:
        self._words_per_batch = words_per_batch
        self._max_paras_per_batch = max_paras_per_batch

    def batch(self, all_paragraphs: list) -> list:
        """Group paragraphs into batches by word count."""
        batches = []
        current_batch = []
        current_words = 0
        for p in all_paragraphs:
            p_words = len(p["full_paragraph"].split())
            if current_batch and (
                current_words + p_words > self._words_per_batch
                or len(current_batch) >= self._max_paras_per_batch
            ):
                batches.append(current_batch)
                current_batch = []
                current_words = 0
            current_batch.append(p)
            current_words += p_words
        if current_batch:
            batches.append(current_batch)
        return batches


class _WordSampler:
    """Pure logic: sample the first n words from extracted batches."""

    def sample(self, batches: list, n: int) -> str:
        """Return the first n words from batch data (for language detection)."""
        words = []
        for batch in batches:
            for p in batch:
                words.extend(p["full_paragraph"].split())
                if len(words) >= n:
                    return " ".join(words[:n])
        return " ".join(words)


class _DocumentExtractor:
    """I/O logic: read a DOCX or PPTX file, extract paragraphs, populate _STATE."""

    def __init__(self, batcher: _ParagraphBatcher) -> None:
        self._batcher = batcher

    def extract(self, file_path: str) -> tuple:
        """Detect format, extract, batch. Returns (batches, stats)."""
        ext = os.path.splitext(file_path)[1].lower()

        if ext == _DOCX_EXTENSION:
            _STATE.reset(DocFormat.DOCX.value, file_path)
            all_paragraphs = self._extract_docx(file_path)
        elif ext == _PPTX_EXTENSION:
            _STATE.reset(DocFormat.PPTX.value, file_path)
            all_paragraphs = self._extract_pptx(file_path)
        else:
            raise ValueError(f"Unsupported file format: {ext}")

        _STATE.batches = self._batcher.batch(all_paragraphs)
        batch_count = len(_STATE.batches)
        _STATE.stats["batch_count"] = batch_count

        _LOGGER.info(f"Extracted: {_STATE.stats['total_paragraphs']} paragraphs")
        stat_parts = [
            f"{k}: {v}"
            for k, v in _STATE.stats.items()
            if k != "total_paragraphs" and k != "batch_count"
        ]
        _LOGGER.info(f"  {' | '.join(stat_parts)}")
        _LOGGER.info(f"  Batches: {batch_count}")

        return _STATE.batches, _STATE.stats

    # --- DOCX extraction ---------------------------------------------------

    def _get_run_text(self, r_elem: Any) -> str:
        """Get text from a run element."""
        return "".join(t.text or "" for t in r_elem.findall(qn("w:t")))

    def _get_formatting_key(self, r_elem: Any) -> str:
        """Get a hashable key representing a run's formatting (canonical rPr XML)."""
        rpr = r_elem.find(qn("w:rPr"))
        if rpr is None:
            return ""
        return etree.tostring(rpr, method="c14n").decode()

    def _merge_docx_runs(self, runs: list) -> list:
        """Merge adjacent runs with identical formatting. Returns [_MergedRun]."""
        if not runs:
            return []
        merged = []
        current_text = self._get_run_text(runs[0])
        current_key = self._get_formatting_key(runs[0])
        current_indices = [0]
        for i in range(1, len(runs)):
            key = self._get_formatting_key(runs[i])
            if key == current_key:
                current_text += self._get_run_text(runs[i])
                current_indices.append(i)
            else:
                merged.append(_MergedRun(current_text, current_indices[:]))
                current_text = self._get_run_text(runs[i])
                current_key = key
                current_indices = [i]
        merged.append(_MergedRun(current_text, current_indices[:]))
        return merged

    def _process_docx_paragraph(
        self, paragraph: Any, para_id: int, all_paragraphs: list
    ) -> int:
        """Process a single DOCX paragraph into extraction format."""
        runs = _get_all_runs_in_order(paragraph)
        if not runs:
            return para_id
        merged = self._merge_docx_runs(runs)
        full_text = "".join(m.text for m in merged)
        if not full_text.strip():
            return para_id
        run_list = []
        merged_run_map = {}
        for run_id, m in enumerate(merged):
            run_list.append({"id": run_id, "text": m.text})
            merged_run_map[run_id] = m.original_indices
        _STATE.run_map[para_id] = {
            "merged_run_map": merged_run_map,
            "run_count": len(runs),
        }
        all_paragraphs.append(
            {"para_id": para_id, "full_paragraph": full_text, "runs": run_list}
        )
        return para_id + 1

    def _extract_docx(self, file_path: str) -> list:
        """Walk DOCX body, tables, headers/footers. Returns all_paragraphs."""
        doc = Document(file_path)
        all_paragraphs = []
        para_id = 0

        for paragraph in doc.paragraphs:
            if (
                paragraph.text
                and paragraph.text.strip()
                and _get_all_runs_in_order(paragraph)
            ):
                para_id = self._process_docx_paragraph(
                    paragraph, para_id, all_paragraphs
                )
            elif paragraph.text and paragraph.text.strip():
                para_id += 1

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        if (
                            paragraph.text
                            and paragraph.text.strip()
                            and _get_all_runs_in_order(paragraph)
                        ):
                            para_id = self._process_docx_paragraph(
                                paragraph, para_id, all_paragraphs
                            )
                        elif paragraph.text and paragraph.text.strip():
                            para_id += 1

        for section in doc.sections:
            for header_footer in [
                section.header,
                section.footer,
                section.first_page_header,
                section.first_page_footer,
            ]:
                if header_footer and header_footer.is_linked_to_previous is False:
                    for paragraph in header_footer.paragraphs:
                        if (
                            paragraph.text
                            and paragraph.text.strip()
                            and _get_all_runs_in_order(paragraph)
                        ):
                            para_id = self._process_docx_paragraph(
                                paragraph, para_id, all_paragraphs
                            )
                        elif paragraph.text and paragraph.text.strip():
                            para_id += 1

        _STATE.stats = {
            "total_paragraphs": len(all_paragraphs),
        }
        return all_paragraphs

    # --- PPTX extraction ---------------------------------------------------

    def _get_pptx_font_key(self, run: Any) -> tuple:
        """Get a hashable key representing a PPTX run's formatting."""
        font = run.font
        try:
            color_val = str(font.color.rgb) if font.color and font.color.rgb else None
        except (AttributeError, TypeError):
            color_val = None
        return (
            str(font.name),
            str(font.size),
            str(font.bold),
            str(font.italic),
            str(font.underline),
            str(color_val),
        )

    def _merge_pptx_runs(self, runs: list) -> list:
        """Merge adjacent PPTX runs with identical font properties. Returns [_MergedRun]."""
        if not runs:
            return []
        merged = []
        current_text = runs[0].text or ""
        current_font = self._get_pptx_font_key(runs[0])
        current_indices = [0]
        for i in range(1, len(runs)):
            font_key = self._get_pptx_font_key(runs[i])
            if font_key == current_font:
                current_text += runs[i].text or ""
                current_indices.append(i)
            else:
                merged.append(_MergedRun(current_text, current_indices[:]))
                current_text = runs[i].text or ""
                current_font = font_key
                current_indices = [i]
        merged.append(_MergedRun(current_text, current_indices[:]))
        return merged

    def _process_pptx_text_frame(
        self, text_frame: Any, para_id: int, all_paragraphs: list
    ) -> int:
        """Process all paragraphs in a PPTX text frame."""
        for paragraph in text_frame.paragraphs:
            if not paragraph.text.strip():
                para_id += 1
                continue
            runs = paragraph.runs
            if not runs:
                all_paragraphs.append(
                    {
                        "para_id": para_id,
                        "full_paragraph": paragraph.text,
                        "runs": [{"id": 0, "text": paragraph.text}],
                    }
                )
                _STATE.run_map[para_id] = {"merged_run_map": {0: [0]}, "run_count": 1}
                para_id += 1
                continue
            merged = self._merge_pptx_runs(runs)
            full_text = "".join(m.text for m in merged)
            if not full_text.strip():
                para_id += 1
                continue
            run_list = []
            merged_run_map = {}
            for run_id, m in enumerate(merged):
                run_list.append({"id": run_id, "text": m.text})
                merged_run_map[run_id] = m.original_indices
            _STATE.run_map[para_id] = {
                "merged_run_map": merged_run_map,
                "run_count": len(runs),
            }
            all_paragraphs.append(
                {"para_id": para_id, "full_paragraph": full_text, "runs": run_list}
            )
            para_id += 1
        return para_id

    def _iter_pptx_shapes(self, shapes: Any) -> Any:
        """Recursively yield all shapes, walking into GroupShape containers."""
        for shape in shapes:
            if shape.shape_type == _PPTX_GROUP_SHAPE_TYPE:
                try:
                    for child in shape.shapes:
                        yield from self._iter_pptx_shapes([child])
                except AttributeError:
                    pass
            else:
                yield shape

    def _extract_pptx(self, file_path: str) -> list:
        """Walk PPTX slides, shapes, tables, notes. Returns all_paragraphs."""
        prs = Presentation(file_path)
        all_paragraphs = []
        para_id = 0
        slide_count = 0

        for slide in prs.slides:
            slide_count += 1
            for shape in self._iter_pptx_shapes(slide.shapes):
                if shape.has_text_frame:
                    para_id = self._process_pptx_text_frame(
                        shape.text_frame, para_id, all_paragraphs
                    )
                if shape.has_table:
                    for row in shape.table.rows:
                        for cell in row.cells:
                            if cell.text_frame:
                                para_id = self._process_pptx_text_frame(
                                    cell.text_frame, para_id, all_paragraphs
                                )
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                para_id = self._process_pptx_text_frame(
                    slide.notes_slide.notes_text_frame, para_id, all_paragraphs
                )

        _STATE.stats = {
            "total_paragraphs": len(all_paragraphs),
            "slides": slide_count,
        }
        return all_paragraphs


def get_first_words(n: int = _DEFAULT_SAMPLE_WORDS) -> str:
    """Return the first n words from batch data (for language detection)."""
    return _WordSampler().sample(_STATE.batches, n)


def extract_document(file_path: str) -> tuple:
    """PUBLIC. Detect format, extract, batch. Returns (batches, stats)."""
    batcher = _ParagraphBatcher(WORDS_PER_BATCH, MAX_PARAS_PER_BATCH)
    extractor = _DocumentExtractor(batcher)
    return extractor.extract(file_path)
