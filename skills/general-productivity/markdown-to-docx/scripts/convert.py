"""
description: Converts a markdown file to a formatted Word document using python-docx, handling headings, inline formatting, lists, tables, images, and other elements.
last_updated: 2026-09-13
origin: original

Markdown to DOCX converter.

Converts a markdown file to a properly formatted Word document using python-docx.
Handles headings, inline formatting (bold, italic, code, hyperlinks), lists (bulleted
and numbered with nesting), horizontal rules, blockquotes, tables, images, and
paragraph continuation.

Dependencies: python-docx, lxml, Pillow (all pre-installed in Amazon Quick sandbox).
"""

import argparse
import logging
import os
import re
import sys
from dataclasses import dataclass, field
from enum import StrEnum
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from lxml import etree


class _PrintStream:
    """Stdout adapter so the logging module can emit human-readable output.

    This is the only place a bare print() is permitted. It forwards each
    formatted log record to the real interpreter stdout.
    """

    def write(self, message: str) -> int:
        text = message.rstrip("\n")
        if text:
            print(text, file=sys.__stdout__)
        return len(message)

    def flush(self) -> None:
        sys.__stdout__.flush()


logging.basicConfig(
    stream=_PrintStream(),
    force=True,
    level=logging.INFO,
    format="%(asctime)s %(levelname)s %(message)s",
)
logger = logging.getLogger(__name__)


# XML namespaces used by the WordprocessingML document format.
WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
RELATIONSHIP_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
XML_NS = "http://www.w3.org/XML/1998/namespace"
HYPERLINK_REL_TYPE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"
)

# Fonts, colors, and sizing used when styling the document.
NORMAL_FONT = "Calibri"
HEADING_FONT = "Calibri Light"
CODE_FONT = "Consolas"
NORMAL_FONT_SIZE_PT = 11
INLINE_CODE_SIZE_PT = 10
CODE_BLOCK_SIZE_PT = 9
NORMAL_LINE_SPACING = 1.15
HYPERLINK_COLOR = "0563C1"
BLOCKQUOTE_BORDER_COLOR = "808080"
TABLE_STYLE = "Table Grid"

# Layout measurements (inches).
IMAGE_WIDTH_INCHES = 5
CODE_BLOCK_INDENT_INCHES = 0.25
BLOCKQUOTE_INDENT_INCHES = 0.5
LIST_FALLBACK_BASE_INDENT_INCHES = 0.25
LIST_NEST_EXTRA_INDENT_INCHES = 0.25

# Heading configuration: (point size, hex color) for Heading 1 through Heading 4.
HEADING_STYLE_SPECS = [
    (24, "1F3864"),
    (18, "2E75B6"),
    (14, "2E75B6"),
    (12, "2E75B6"),
]
MAX_HEADING_LEVEL_INDEX = 3

# Markdown parsing thresholds and placeholders.
NESTED_LIST_INDENT_THRESHOLD = 4
SEPARATOR_CELL_CHARS = set("-: ")
TABLE_SEP_PLACEHOLDER_WIDTH = 10


class InlineType(StrEnum):
    """The finite set of inline markdown constructs recognized by the parser."""

    LINK = "link"
    BOLD_ITALIC = "bold_italic"
    BOLD = "bold"
    ITALIC = "italic"
    CODE = "code"


class LineType(StrEnum):
    """The finite set of block-level line classifications."""

    CODE_FENCE = "code_fence"
    HEADING = "heading"
    HR = "hr"
    IMAGE = "image"
    BLOCKQUOTE = "blockquote"
    TABLE_SEP = "table_sep"
    TABLE_ROW = "table_row"
    NUMBERED = "numbered"
    BULLET = "bullet"
    BLANK = "blank"
    TEXT = "text"


# Ordered inline patterns; the earliest match in the text wins.
INLINE_PATTERNS = [
    (InlineType.LINK, re.compile(r"\[([^\]]+)\]\(([^)]+)\)")),
    (InlineType.BOLD_ITALIC, re.compile(r"\*\*\*(.+?)\*\*\*")),
    (InlineType.BOLD, re.compile(r"\*\*(.+?)\*\*")),
    (InlineType.ITALIC, re.compile(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)")),
    (InlineType.CODE, re.compile(r"`([^`]+)`")),
]

HEADING_RE = re.compile(r"^(#{1,6})\s+(.+)$")
BULLET_RE = re.compile(r"^(\s*)([-*+])\s+(.+)$")
NUMBERED_RE = re.compile(r"^(\s*)(\d+)\.\s+(.+)$")
BLOCKQUOTE_RE = re.compile(r"^>\s*(.*)$")
HR_RE = re.compile(r"^(---+|\*\*\*+|___+)\s*$")
TABLE_ROW_RE = re.compile(r"^\|(.+)\|$")
TABLE_SEP_RE = re.compile(r"^\|[\s:]*[-]+[\s:]*")
IMAGE_RE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)$")
CODE_FENCE_RE = re.compile(r"^```")


@dataclass(frozen=True)
class InlineSegment:
    """A run of inline text with its formatting attributes."""

    text: str
    bold: bool = False
    italic: bool = False
    code: bool = False
    link: str | None = None


@dataclass(frozen=True)
class ClassifiedLine:
    """The block-level classification of a single markdown line."""

    line_type: LineType
    text: str = ""
    level: int = 0
    indent: int = 0
    alt: str = ""
    path: str = ""
    cells: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class HeadingElement:
    """A heading paragraph at a resolved python-docx level."""

    doc_level: int
    segments: list[InlineSegment] = field(default_factory=list)


@dataclass(frozen=True)
class HorizontalRuleElement:
    """A horizontal rule rendered as a bottom-bordered paragraph."""


@dataclass(frozen=True)
class BlockquoteElement:
    """A blockquote paragraph with a left border and italic runs."""

    segments: list[InlineSegment] = field(default_factory=list)


@dataclass(frozen=True)
class BulletElement:
    """A bulleted list item at a given indentation."""

    indent: int
    segments: list[InlineSegment] = field(default_factory=list)


@dataclass(frozen=True)
class NumberedElement:
    """A numbered list item at a given indentation."""

    indent: int
    segments: list[InlineSegment] = field(default_factory=list)


@dataclass(frozen=True)
class ImageElement:
    """An image reference with alt text and an unresolved source path."""

    alt: str
    path: str


@dataclass(frozen=True)
class ParagraphElement:
    """A body paragraph assembled from continuation lines."""

    segments: list[InlineSegment] = field(default_factory=list)


@dataclass(frozen=True)
class TableCell:
    """A single table cell holding its parsed inline segments."""

    segments: list[InlineSegment] = field(default_factory=list)


@dataclass(frozen=True)
class TableRow:
    """A single table row holding its cells."""

    cells: list[TableCell] = field(default_factory=list)


@dataclass(frozen=True)
class TableElement:
    """A table composed of data rows (separator rows already removed)."""

    rows: list[TableRow] = field(default_factory=list)


@dataclass(frozen=True)
class CodeBlockElement:
    """A fenced code block rendered as a monospace, indented paragraph."""

    text: str


DocumentElement = (
    HeadingElement
    | HorizontalRuleElement
    | BlockquoteElement
    | BulletElement
    | NumberedElement
    | ImageElement
    | ParagraphElement
    | TableElement
    | CodeBlockElement
)


@dataclass(frozen=True)
class ConversionRequest:
    """Inputs for a single markdown to docx conversion."""

    source_path: str
    output_path: str
    title: str | None = None


class InlineFormatParser:
    """Parses a line of text into ordered, formatted inline segments."""

    def parse(self, text: str) -> list[InlineSegment]:
        segments: list[InlineSegment] = []
        self._parse_recursive(text, segments)
        return segments

    def _parse_recursive(self, text: str, segments: list[InlineSegment]) -> None:
        if not text:
            return
        earliest_match = None
        earliest_start = len(text)
        earliest_type: InlineType | None = None
        for ptype, pattern in INLINE_PATTERNS:
            m = pattern.search(text)
            if m and m.start() < earliest_start:
                earliest_match = m
                earliest_start = m.start()
                earliest_type = ptype
        if earliest_match is None:
            segments.append(InlineSegment(text=text))
            return
        before = text[:earliest_start]
        if before:
            segments.append(InlineSegment(text=before))
        segments.append(self._segment_for_match(earliest_type, earliest_match))
        after = text[earliest_match.end() :]
        if after:
            self._parse_recursive(after, segments)

    def _segment_for_match(
        self, inline_type: InlineType | None, match: "re.Match[str]"
    ) -> InlineSegment:
        if inline_type == InlineType.LINK:
            return InlineSegment(text=match.group(1), link=match.group(2))
        if inline_type == InlineType.BOLD_ITALIC:
            return InlineSegment(text=match.group(1), bold=True, italic=True)
        if inline_type == InlineType.BOLD:
            return InlineSegment(text=match.group(1), bold=True)
        if inline_type == InlineType.ITALIC:
            return InlineSegment(text=match.group(1), italic=True)
        return InlineSegment(text=match.group(1), code=True)


class MarkdownDocumentParser:
    """Parses markdown source lines into an ordered list of document elements.

    This is pure logic: it performs no I/O and never touches python-docx. The
    resulting elements describe what to render, in order, so the renderer can
    reproduce them faithfully.
    """

    def __init__(self, inline_parser: InlineFormatParser) -> None:
        self._inline = inline_parser

    def parse(self, lines: list[str], title: str | None) -> list[DocumentElement]:
        elements: list[DocumentElement] = []
        in_code_block = False
        code_lines: list[str] = []
        in_table = False
        table_rows: list[list[str]] = []
        paragraph_buffer: list[str] = []
        first_h1_seen = False

        def flush_paragraph_buffer() -> None:
            nonlocal paragraph_buffer
            if paragraph_buffer:
                combined = " ".join(paragraph_buffer)
                segments = self._inline.parse(combined)
                elements.append(ParagraphElement(segments=segments))
                paragraph_buffer = []

        def flush_table() -> None:
            nonlocal in_table, table_rows
            table_element = self._build_table_element(table_rows)
            if table_element is not None:
                elements.append(table_element)
            in_table = False
            table_rows = []

        def flush_code_block() -> None:
            nonlocal code_lines
            if code_lines:
                elements.append(CodeBlockElement(text="\n".join(code_lines)))
                code_lines = []

        for raw_line in lines:
            line = raw_line.rstrip("\n")

            if in_code_block:
                if CODE_FENCE_RE.match(line):
                    in_code_block = False
                    flush_code_block()
                else:
                    code_lines.append(line)
                continue

            classified = self._classify_line(line)
            line_type = classified.line_type

            if line_type == LineType.CODE_FENCE:
                flush_paragraph_buffer()
                if in_table:
                    flush_table()
                in_code_block = True
                continue

            if line_type == LineType.TABLE_ROW:
                flush_paragraph_buffer()
                if not in_table:
                    in_table = True
                    table_rows = []
                table_rows.append(classified.cells)
                continue
            elif line_type == LineType.TABLE_SEP:
                if in_table:
                    table_rows.append(["---"] * TABLE_SEP_PLACEHOLDER_WIDTH)
                continue
            elif in_table:
                flush_table()

            if line_type == LineType.HEADING:
                flush_paragraph_buffer()
                level = classified.level
                text = classified.text
                doc_level = min(level - 1, MAX_HEADING_LEVEL_INDEX)
                if level == 1 and not first_h1_seen:
                    first_h1_seen = True
                    if title:
                        text = title
                segments = self._inline.parse(text)
                elements.append(HeadingElement(doc_level=doc_level, segments=segments))

            elif line_type == LineType.HR:
                flush_paragraph_buffer()
                elements.append(HorizontalRuleElement())

            elif line_type == LineType.BLOCKQUOTE:
                flush_paragraph_buffer()
                content = classified.text
                segments = self._inline.parse(content) if content else []
                elements.append(BlockquoteElement(segments=segments))

            elif line_type == LineType.BULLET:
                flush_paragraph_buffer()
                segments = self._inline.parse(classified.text)
                elements.append(
                    BulletElement(indent=classified.indent, segments=segments)
                )

            elif line_type == LineType.NUMBERED:
                flush_paragraph_buffer()
                segments = self._inline.parse(classified.text)
                elements.append(
                    NumberedElement(indent=classified.indent, segments=segments)
                )

            elif line_type == LineType.IMAGE:
                flush_paragraph_buffer()
                elements.append(ImageElement(alt=classified.alt, path=classified.path))

            elif line_type == LineType.BLANK:
                flush_paragraph_buffer()

            elif line_type == LineType.TEXT:
                paragraph_buffer.append(classified.text)

        flush_paragraph_buffer()
        if in_table:
            flush_table()
        if in_code_block:
            flush_code_block()

        return elements

    def _classify_line(self, line: str) -> ClassifiedLine:
        m = CODE_FENCE_RE.match(line)
        if m:
            return ClassifiedLine(line_type=LineType.CODE_FENCE, text=line[3:].strip())
        m = HEADING_RE.match(line)
        if m:
            return ClassifiedLine(
                line_type=LineType.HEADING,
                level=len(m.group(1)),
                text=m.group(2).strip(),
            )
        m = HR_RE.match(line)
        if m:
            return ClassifiedLine(line_type=LineType.HR)
        m = IMAGE_RE.match(line)
        if m:
            return ClassifiedLine(
                line_type=LineType.IMAGE, alt=m.group(1), path=m.group(2)
            )
        m = BLOCKQUOTE_RE.match(line)
        if m:
            return ClassifiedLine(line_type=LineType.BLOCKQUOTE, text=m.group(1))
        m = TABLE_SEP_RE.match(line)
        if m:
            return ClassifiedLine(line_type=LineType.TABLE_SEP)
        m = TABLE_ROW_RE.match(line)
        if m:
            return ClassifiedLine(
                line_type=LineType.TABLE_ROW,
                cells=[c.strip() for c in m.group(1).split("|")],
            )
        m = NUMBERED_RE.match(line)
        if m:
            return ClassifiedLine(
                line_type=LineType.NUMBERED,
                indent=len(m.group(1)),
                text=m.group(3),
            )
        m = BULLET_RE.match(line)
        if m:
            return ClassifiedLine(
                line_type=LineType.BULLET,
                indent=len(m.group(1)),
                text=m.group(3),
            )
        if line.strip() == "":
            return ClassifiedLine(line_type=LineType.BLANK)
        return ClassifiedLine(line_type=LineType.TEXT, text=line)

    def _build_table_element(self, table_rows: list[list[str]]) -> TableElement | None:
        if not table_rows:
            return None
        data_rows = [
            r
            for r in table_rows
            if not all(set(cell.strip()) <= SEPARATOR_CELL_CHARS for cell in r)
        ]
        if not data_rows:
            return None
        rows: list[TableRow] = []
        for row_cells in data_rows:
            cells = [
                TableCell(segments=self._inline.parse(cell_text))
                for cell_text in row_cells
            ]
            rows.append(TableRow(cells=cells))
        return TableElement(rows=rows)


class DocxConverter:
    """Renders parsed markdown elements into a Word document on disk.

    This class owns all I/O (reading the source, resolving image paths, and
    saving the .docx). It depends on the pure MarkdownDocumentParser, never the
    reverse.
    """

    def __init__(self, parser: MarkdownDocumentParser) -> None:
        self._parser = parser

    def convert(self, request: ConversionRequest) -> str:
        lines = self._read_lines(request.source_path)
        elements = self._parser.parse(lines, request.title)
        document = Document()
        self._setup_styles(document)
        source_dir = os.path.dirname(os.path.abspath(request.source_path))
        for element in elements:
            self._render_element(document, element, source_dir)
        document.save(request.output_path)
        return request.output_path

    def _read_lines(self, source_path: str) -> list[str]:
        with Path(source_path).open(encoding="utf-8") as handle:
            return handle.readlines()

    def _setup_styles(self, doc: Document) -> None:
        style = doc.styles["Normal"]
        font = style.font
        font.name = NORMAL_FONT
        font.size = Pt(NORMAL_FONT_SIZE_PT)
        pf = style.paragraph_format
        pf.line_spacing = NORMAL_LINE_SPACING
        for level, (size, color_hex) in enumerate(HEADING_STYLE_SPECS):
            style_name = f"Heading {level + 1}"
            try:
                h_style = doc.styles[style_name]
                h_style.font.name = HEADING_FONT
                h_style.font.size = Pt(size)
                h_style.font.bold = True
                r, g, b = (
                    int(color_hex[0:2], 16),
                    int(color_hex[2:4], 16),
                    int(color_hex[4:6], 16),
                )
                h_style.font.color.rgb = RGBColor(r, g, b)
            except KeyError:
                pass

    def _render_element(
        self, doc: Document, element: DocumentElement, source_dir: str
    ) -> None:
        if isinstance(element, HeadingElement):
            self._render_heading(doc, element)
        elif isinstance(element, HorizontalRuleElement):
            self._render_horizontal_rule(doc)
        elif isinstance(element, BlockquoteElement):
            self._render_blockquote(doc, element)
        elif isinstance(element, BulletElement):
            self._render_bullet(doc, element)
        elif isinstance(element, NumberedElement):
            self._render_numbered(doc, element)
        elif isinstance(element, ImageElement):
            self._render_image(doc, element, source_dir)
        elif isinstance(element, ParagraphElement):
            self._render_paragraph(doc, element)
        elif isinstance(element, TableElement):
            self._render_table(doc, element)
        elif isinstance(element, CodeBlockElement):
            self._render_code_block(doc, element)

    def _render_heading(self, doc: Document, element: HeadingElement) -> None:
        p = doc.add_heading("", level=element.doc_level)
        self._apply_segments(p, element.segments)

    def _render_horizontal_rule(self, doc: Document) -> None:
        p = doc.add_paragraph()
        pPr = p._p.find(f"{{{WORD_NS}}}pPr")
        if pPr is None:
            pPr = etree.SubElement(p._p, f"{{{WORD_NS}}}pPr")
            p._p.insert(0, pPr)
        pBdr = etree.SubElement(pPr, f"{{{WORD_NS}}}pBdr")
        bottom = etree.SubElement(pBdr, f"{{{WORD_NS}}}bottom")
        bottom.set(f"{{{WORD_NS}}}val", "single")
        bottom.set(f"{{{WORD_NS}}}sz", "6")
        bottom.set(f"{{{WORD_NS}}}space", "1")
        bottom.set(f"{{{WORD_NS}}}color", "auto")

    def _render_blockquote(self, doc: Document, element: BlockquoteElement) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(BLOCKQUOTE_INDENT_INCHES)
        pPr = p._p.find(f"{{{WORD_NS}}}pPr")
        if pPr is None:
            pPr = etree.SubElement(p._p, f"{{{WORD_NS}}}pPr")
            p._p.insert(0, pPr)
        pBdr = etree.SubElement(pPr, f"{{{WORD_NS}}}pBdr")
        left = etree.SubElement(pBdr, f"{{{WORD_NS}}}left")
        left.set(f"{{{WORD_NS}}}val", "single")
        left.set(f"{{{WORD_NS}}}sz", "12")
        left.set(f"{{{WORD_NS}}}space", "4")
        left.set(f"{{{WORD_NS}}}color", BLOCKQUOTE_BORDER_COLOR)
        if element.segments:
            self._apply_segments(p, element.segments)
            for run in p.runs:
                run.italic = True

    def _render_bullet(self, doc: Document, element: BulletElement) -> None:
        p = self._list_paragraph(doc, element.indent, "List Bullet", "List Bullet 2")
        self._apply_segments(p, element.segments)

    def _render_numbered(self, doc: Document, element: NumberedElement) -> None:
        p = self._list_paragraph(doc, element.indent, "List Number", "List Number 2")
        self._apply_segments(p, element.segments)

    def _list_paragraph(
        self, doc: Document, indent: int, base_style: str, nested_style: str
    ):
        is_nested = indent >= NESTED_LIST_INDENT_THRESHOLD
        style_name = nested_style if is_nested else base_style
        try:
            return doc.add_paragraph(style=style_name)
        except KeyError:
            p = doc.add_paragraph()
            extra = LIST_NEST_EXTRA_INDENT_INCHES if is_nested else 0
            p.paragraph_format.left_indent = Inches(
                LIST_FALLBACK_BASE_INDENT_INCHES + extra
            )
            return p

    def _render_image(
        self, doc: Document, element: ImageElement, source_dir: str
    ) -> None:
        img_path = element.path
        if not os.path.isabs(img_path):
            img_path = os.path.join(source_dir, img_path)
        if os.path.exists(img_path):
            try:
                doc.add_picture(img_path, width=Inches(IMAGE_WIDTH_INCHES))
            except Exception:
                p = doc.add_paragraph()
                run = p.add_run(f"[Image: {element.alt or img_path}]")
                run.italic = True
        else:
            p = doc.add_paragraph()
            run = p.add_run(f"[Image not found: {element.alt or img_path}]")
            run.italic = True

    def _render_paragraph(self, doc: Document, element: ParagraphElement) -> None:
        p = doc.add_paragraph()
        self._apply_segments(p, element.segments)

    def _render_table(self, doc: Document, element: TableElement) -> None:
        data_rows = element.rows
        num_cols = max(len(row.cells) for row in data_rows)
        tbl = doc.add_table(rows=len(data_rows), cols=num_cols)
        tbl.style = TABLE_STYLE
        for i, row in enumerate(data_rows):
            for j, cell in enumerate(row.cells):
                if j < num_cols:
                    doc_cell = tbl.rows[i].cells[j]
                    doc_cell.text = ""
                    p = doc_cell.paragraphs[0]
                    self._apply_segments(p, cell.segments)
        for doc_cell in tbl.rows[0].cells:
            for paragraph in doc_cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True

    def _render_code_block(self, doc: Document, element: CodeBlockElement) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(CODE_BLOCK_INDENT_INCHES)
        run = p.add_run(element.text)
        run.font.name = CODE_FONT
        run.font.size = Pt(CODE_BLOCK_SIZE_PT)

    def _apply_segments(self, paragraph, segments: list[InlineSegment]) -> None:
        for seg in segments:
            if seg.link:
                self._add_hyperlink(paragraph, seg.text, seg.link)
            elif seg.code:
                run = paragraph.add_run(seg.text)
                run.font.name = CODE_FONT
                run.font.size = Pt(INLINE_CODE_SIZE_PT)
            else:
                run = paragraph.add_run(seg.text)
                if seg.bold:
                    run.bold = True
                if seg.italic:
                    run.italic = True

    def _add_hyperlink(self, paragraph, text: str, url: str) -> None:
        part = paragraph.part
        r_id = part.relate_to(url, HYPERLINK_REL_TYPE, is_external=True)

        hyperlink = etree.SubElement(paragraph._p, f"{{{WORD_NS}}}hyperlink")
        hyperlink.set(f"{{{RELATIONSHIP_NS}}}id", r_id)

        run_elem = etree.SubElement(hyperlink, f"{{{WORD_NS}}}r")
        rPr = etree.SubElement(run_elem, f"{{{WORD_NS}}}rPr")

        color = etree.SubElement(rPr, f"{{{WORD_NS}}}color")
        color.set(f"{{{WORD_NS}}}val", HYPERLINK_COLOR)

        u = etree.SubElement(rPr, f"{{{WORD_NS}}}u")
        u.set(f"{{{WORD_NS}}}val", "single")

        t_elem = etree.SubElement(run_elem, f"{{{WORD_NS}}}t")
        t_elem.set(f"{{{XML_NS}}}space", "preserve")
        t_elem.text = text


def convert_markdown_to_docx(
    source_path: str, output_path: str, title: str | None = None
) -> str:
    """
    Convert a markdown file to a formatted Word document.

    Args:
        source_path: Path to the .md file
        output_path: Path for the output .docx file
        title: Optional document title (overrides the first H1 if provided)

    Returns:
        Path to the generated .docx file
    """
    inline_parser = InlineFormatParser()
    document_parser = MarkdownDocumentParser(inline_parser)
    converter = DocxConverter(document_parser)
    request = ConversionRequest(
        source_path=source_path, output_path=output_path, title=title
    )
    return converter.convert(request)


def _build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Convert a markdown file to a formatted Word document."
    )
    parser.add_argument("input", help="Path to the input markdown (.md) file.")
    parser.add_argument("output", help="Path for the output .docx file.")
    parser.add_argument(
        "title",
        nargs="?",
        default=None,
        help="Optional document title that overrides the first H1 heading.",
    )
    return parser


def main() -> int:
    parser = _build_arg_parser()
    args = parser.parse_args()
    result = convert_markdown_to_docx(args.input, args.output, args.title)
    logger.info("Converted: %s", result)
    return 0


if __name__ == "__main__":
    sys.exit(main())
