"""
description: Create confined Word and Markdown petrophysical reports from evaluated well-log results.
last_updated: 2026-09-16
origin: original
"""

from __future__ import annotations

import argparse
import json
import logging
import re
import sys
from pathlib import Path
from typing import Any

WELL_ID_PATTERN = re.compile(r"^[A-Za-z0-9._-]+$")
DEFAULT_WELL_ID = "SYN-LAS-1"
DEFAULT_DATA_ROOT = "automated-log-analysis-data"
PARAMETER_LABELS = {
    "area_acres": "Drainage area (acres)",
    "matrix_density": "Matrix density (g/cc)",
    "fluid_density": "Fluid density (g/cc)",
    "a": "Archie tortuosity factor (a)",
    "m": "Archie cementation exponent (m)",
    "n": "Archie saturation exponent (n)",
    "rw": "Formation water resistivity Rw (ohm-m)",
    "phie_min": "Net pay PHIE minimum",
    "sw_max": "Net pay Sw maximum",
    "vsh_max": "Net pay Vshale maximum",
    "boi": "Oil formation volume factor Boi (RB/STB)",
    "rf": "Recovery factor",
}
ZONE_COLUMNS = (
    ("zone_id", "Zone"),
    ("top", "Top (ft)"),
    ("base", "Base (ft)"),
    ("thickness", "Thickness (ft)"),
    ("avg_phie", "Avg PHIE"),
    ("avg_sw", "Avg Sw"),
    ("avg_rt", "Avg RT (ohm-m)"),
)


def parameter_label(key: str) -> str:
    """Return the human-readable label for an evaluation parameter key."""
    return PARAMETER_LABELS.get(key, key.replace("_", " ").capitalize())


SW_UNAVAILABLE = "n/a (no resistivity)"


def format_optional(value: Any) -> str:
    """Render an optional metric, never the literal None."""
    return "n/a" if value is None else str(value)


def format_sw(value: Any) -> str:
    """Render a water-saturation value, naming the reason when it is null."""
    return SW_UNAVAILABLE if value is None else str(value)


def format_stb(value: Any, note: str | None) -> str:
    """Render an OOIP-family volume, or the evaluator's note when it is null."""
    if value is None:
        return f"not estimated ({note})" if note else "not estimated"
    return f"{value} STB"


def net_pay_basis(net_pay: dict[str, Any]) -> str:
    """Describe what screened the net pay so limited accuracy is explicit."""
    accuracy = str(net_pay.get("accuracy", "full"))
    if net_pay.get("sw_cutoff_applied", True):
        return "porosity, shale, and water-saturation cutoffs (full accuracy)"
    return f"porosity and shale cutoffs only; limited accuracy ({accuracy})"


class _PrintStream:
    def write(self, message: str) -> None:
        if message.strip():
            print(message, end="")

    def flush(self) -> None:
        pass


logging.basicConfig(
    level=logging.INFO, format="%(message)s", stream=_PrintStream(), force=True
)
logger = logging.getLogger(__name__)


class WorkspacePathResolver:
    """Resolve one relative path beneath an explicit workspace directory."""

    def __init__(self, workspace_dir: str, relative_path: str, label: str) -> None:
        self._workspace_dir = workspace_dir
        self._relative_path = relative_path
        self._label = label

    def resolve(self) -> Path:
        if not self._workspace_dir:
            raise ValueError("workspace_dir is required")
        requested = Path(self._relative_path)
        if requested.is_absolute():
            raise ValueError(f"{self._label} must be relative to workspace_dir")
        if ".." in requested.parts:
            raise ValueError(f"{self._label} must not contain parent traversal")
        workspace = Path(self._workspace_dir).resolve()
        resolved = (workspace / requested).resolve()
        if resolved != workspace and workspace not in resolved.parents:
            raise ValueError(f"{self._label} escapes workspace_dir")
        return resolved


class MarkdownReportBuilder:
    """Build one deterministic Markdown report."""

    def __init__(
        self,
        parsed: dict[str, Any],
        evaluation: dict[str, Any],
        recommendations: list[str],
        plot_path: Path | None,
    ) -> None:
        self._parsed = parsed
        self._evaluation = evaluation
        self._recommendations = recommendations
        self._plot_path = plot_path

    def build(self) -> str:
        well_id = str(self._evaluation.get("well_id", DEFAULT_WELL_ID))
        depth = self._parsed.get("depth_range", {})
        parameters = self._evaluation.get("parameters", {})
        net_pay = self._evaluation.get("net_pay_summary", {})
        ooip = self._evaluation.get("ooip", {})
        qc = self._evaluation.get("qc_summary", {})
        zones = self._evaluation.get("zones", [])
        lines = [
            f"# Automated Log Analysis — {well_id}",
            "",
            "## Well information",
            "",
            f"- Well ID: {well_id}",
            f"- Depth range: {depth.get('minimum')}–{depth.get('maximum')} {depth.get('unit', '')}",
            f"- Samples: {self._parsed.get('sample_count', 0)}",
            "",
            "## Methodology",
            "",
            "Linear gamma-ray shale volume, density porosity, Archie water saturation, shale-corrected effective porosity, cutoff-based net pay, and volumetric OOIP.",
            "",
            "### Parameters",
            "",
            "| Parameter | Value |",
            "| --- | ---: |",
        ]
        lines.extend(
            f"| {parameter_label(name)} | {value} |"
            for name, value in parameters.items()
        )
        lines.extend(
            [
                "",
                "## Cutoffs",
                "",
                "| Cutoff | Value |",
                "| --- | ---: |",
                f"| PHIE minimum | {parameters.get('phie_min')} |",
                f"| SW maximum | {parameters.get('sw_max')} |",
                f"| Vshale maximum | {parameters.get('vsh_max')} |",
                "",
                "## Composite log",
                "",
                f"![Composite log]({self._plot_path.name})"
                if self._plot_path is not None
                else "Composite log not supplied.",
                "",
                "## Net pay zones",
                "",
                "| Zone | Top (ft) | Base (ft) | Thickness (ft) | Avg PHIE | Avg SW | Avg RT |",
                "| ---: | ---: | ---: | ---: | ---: | ---: | ---: |",
            ]
        )
        for zone in zones:
            lines.append(
                f"| {zone.get('zone_id')} | {zone.get('top')} | {zone.get('base')} | "
                f"{zone.get('thickness')} | {zone.get('avg_phie')} | {format_sw(zone.get('avg_sw'))} | "
                f"{format_optional(zone.get('avg_rt'))} |"
            )
        lines.extend(
            [
                "",
                "## Summary metrics",
                "",
                f"- Net pay: {net_pay.get('net_pay_ft')} ft",
                f"- Net pay basis: {net_pay_basis(net_pay)}",
                f"- Gross interval: {net_pay.get('gross_ft')} ft",
                f"- Net-to-gross: {net_pay.get('net_to_gross')}",
                f"- Average pay PHIE: {net_pay.get('avg_phie_pay')}",
                f"- Average pay SW: {format_sw(net_pay.get('avg_sw_pay'))}",
                "",
                "## OOIP estimate",
                "",
                f"- OOIP: {format_stb(ooip.get('ooip_stb'), ooip.get('note'))}",
                f"- Recoverable: {format_stb(ooip.get('recoverable_stb'), ooip.get('note'))}",
                "",
                "## QC notes",
                "",
                f"- Overall quality: {qc.get('overall_quality')}",
            ]
        )
        qc_warnings = qc.get("warnings", [])
        evaluation_warnings = self._evaluation.get("warnings", [])
        for warning in [*qc_warnings, *evaluation_warnings]:
            lines.append(f"- {warning}")
        lines.extend(["", "## Recommendations", ""])
        if self._recommendations:
            lines.extend(f"- {item}" for item in self._recommendations)
        else:
            lines.append("- None supplied.")
        return "\n".join(lines) + "\n"


class WordReportWriter:
    """Write one Word report when python-docx is available."""

    def __init__(
        self,
        output_path: Path,
        parsed: dict[str, Any],
        evaluation: dict[str, Any],
        recommendations: list[str],
        plot_path: Path | None,
    ) -> None:
        self._output_path = output_path
        self._parsed = parsed
        self._evaluation = evaluation
        self._recommendations = recommendations
        self._plot_path = plot_path

    def write(self) -> str | None:
        try:
            from docx import Document
            from docx.shared import Inches
        except ImportError:
            return None
        document = Document()
        well_id = str(self._evaluation.get("well_id", DEFAULT_WELL_ID))
        document.add_heading(f"Automated Log Analysis — {well_id}", 0)
        self._add_well_information(document)
        self._add_methodology(document)
        self._add_cutoffs(document)
        if self._plot_path is not None:
            document.add_heading("Composite log", level=1)
            document.add_picture(str(self._plot_path), width=Inches(6.5))
        self._add_zones(document)
        self._add_summary(document)
        self._add_qc(document)
        document.add_heading("Recommendations", level=1)
        if self._recommendations:
            for recommendation in self._recommendations:
                document.add_paragraph(recommendation, style="List Bullet")
        else:
            document.add_paragraph("None supplied.")
        document.save(self._output_path)
        return str(self._output_path)

    def _add_well_information(self, document: Any) -> None:
        document.add_heading("Well information", level=1)
        depth = self._parsed.get("depth_range", {})
        rows = [
            ("Well ID", self._evaluation.get("well_id")),
            ("Start depth", depth.get("minimum")),
            ("Stop depth", depth.get("maximum")),
            ("Depth unit", depth.get("unit")),
            ("Samples", self._parsed.get("sample_count")),
        ]
        self._add_key_value_table(document, rows)

    def _add_methodology(self, document: Any) -> None:
        document.add_heading("Methodology", level=1)
        document.add_paragraph(
            "Linear gamma-ray shale volume, density porosity, Archie water saturation, shale-corrected effective porosity, cutoff-based net pay, and volumetric OOIP."
        )
        rows = [
            (parameter_label(key), value)
            for key, value in self._evaluation.get("parameters", {}).items()
        ]
        self._add_key_value_table(document, rows, "Parameter")

    def _add_cutoffs(self, document: Any) -> None:
        document.add_heading("Cutoffs", level=1)
        parameters = self._evaluation.get("parameters", {})
        self._add_key_value_table(
            document,
            [
                ("PHIE minimum", parameters.get("phie_min")),
                ("SW maximum", parameters.get("sw_max")),
                ("Vshale maximum", parameters.get("vsh_max")),
            ],
            "Cutoff",
        )

    def _add_zones(self, document: Any) -> None:
        document.add_heading("Net pay zones", level=1)
        table = document.add_table(rows=1, cols=len(ZONE_COLUMNS))
        for index, (_key, label) in enumerate(ZONE_COLUMNS):
            table.rows[0].cells[index].text = label
        for zone in self._evaluation.get("zones", []):
            cells = table.add_row().cells
            for index, (key, _label) in enumerate(ZONE_COLUMNS):
                value = zone.get(key, "")
                if key == "avg_sw":
                    cells[index].text = format_sw(value)
                else:
                    cells[index].text = format_optional(value)

    def _add_summary(self, document: Any) -> None:
        document.add_heading("Summary metrics", level=1)
        net_pay = self._evaluation.get("net_pay_summary", {})
        ooip = self._evaluation.get("ooip", {})
        self._add_key_value_table(
            document,
            [
                ("Net pay (ft)", net_pay.get("net_pay_ft")),
                ("Net pay basis", net_pay_basis(net_pay)),
                ("Gross interval (ft)", net_pay.get("gross_ft")),
                ("Net-to-gross", net_pay.get("net_to_gross")),
                ("Average pay PHIE", net_pay.get("avg_phie_pay")),
                ("Average pay SW", format_sw(net_pay.get("avg_sw_pay"))),
                ("OOIP", format_stb(ooip.get("ooip_stb"), ooip.get("note"))),
                (
                    "Recoverable",
                    format_stb(ooip.get("recoverable_stb"), ooip.get("note")),
                ),
            ],
            "Metric",
        )

    def _add_qc(self, document: Any) -> None:
        document.add_heading("QC notes", level=1)
        qc = self._evaluation.get("qc_summary", {})
        document.add_paragraph(f"Overall quality: {qc.get('overall_quality')}")
        for warning in [*qc.get("warnings", []), *self._evaluation.get("warnings", [])]:
            document.add_paragraph(str(warning), style="List Bullet")

    @staticmethod
    def _add_key_value_table(
        document: Any, rows: list[tuple[str, Any]], header: str = "Item"
    ) -> None:
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = header
        table.rows[0].cells[1].text = "Value"
        for key, value in rows:
            cells = table.add_row().cells
            cells[0].text = str(key)
            cells[1].text = str(value)


def run(request: dict[str, Any]) -> dict[str, Any]:
    """Create Word and Markdown reports from an already-parsed request object."""
    try:
        if not isinstance(request, dict):
            raise ValueError("input must be a JSON object")
        workspace_dir = request.get("workspace_dir")
        if not isinstance(workspace_dir, str) or not workspace_dir:
            raise ValueError("workspace_dir is required")
        well_id = str(request.get("well_id", DEFAULT_WELL_ID))
        if not WELL_ID_PATTERN.fullmatch(well_id):
            raise ValueError("well_id must match ^[A-Za-z0-9._-]+$")
        data_root = str(request.get("data_root", DEFAULT_DATA_ROOT))
        output_root = WorkspacePathResolver(
            workspace_dir, data_root, "data_root"
        ).resolve()
        parsed_relative = str(
            request.get("parsed_path", str(Path(data_root) / f"{well_id}_parsed.json"))
        )
        evaluation_relative = str(
            request.get(
                "evaluation_path", str(Path(data_root) / f"{well_id}_evaluation.json")
            )
        )
        parsed_path = WorkspacePathResolver(
            workspace_dir, parsed_relative, "parsed_path"
        ).resolve()
        evaluation_path = WorkspacePathResolver(
            workspace_dir, evaluation_relative, "evaluation_path"
        ).resolve()
        parsed = json.loads(parsed_path.read_text(encoding="utf-8"))
        evaluation = json.loads(evaluation_path.read_text(encoding="utf-8"))
        recommendations = request.get("recommendations", [])
        if not isinstance(recommendations, list) or any(
            not isinstance(item, str) for item in recommendations
        ):
            raise ValueError("recommendations must be an array of strings")
        plot_relative = request.get(
            "plot_path", str(Path(data_root) / f"{well_id}_log_plot.png")
        )
        plot_path = WorkspacePathResolver(
            workspace_dir, str(plot_relative), "plot_path"
        ).resolve()
        if not plot_path.is_file():
            plot_path = None
        output_root.mkdir(parents=True, exist_ok=True)
        markdown_path = output_root / f"{well_id}_report.md"
        word_path = output_root / f"{well_id}_report.docx"
        markdown = MarkdownReportBuilder(
            parsed, evaluation, recommendations, plot_path
        ).build()
        markdown_path.write_text(markdown, encoding="utf-8", newline="\n")
        word_result = WordReportWriter(
            word_path, parsed, evaluation, recommendations, plot_path
        ).write()
        warnings = (
            []
            if word_result is not None
            else ["python-docx unavailable; Word report skipped"]
        )
        return {
            "status": "created",
            "well_id": well_id,
            "docx_path": word_result,
            "markdown_path": str(markdown_path),
            "warnings": warnings,
        }
    except (ValueError, OSError, json.JSONDecodeError, TypeError) as error:
        return {"status": "error", "error": str(error)}


def main() -> int:
    parser = argparse.ArgumentParser(description="Create Word and Markdown reports.")
    parser.add_argument("--input", required=True, help="JSON request object")
    args = parser.parse_args()
    try:
        request = json.loads(args.input)
    except json.JSONDecodeError as error:
        logger.info(json.dumps({"status": "error", "error": str(error)}))
        return 2
    result = run(request)
    logger.info(json.dumps(result, separators=(",", ":")))
    return 0 if result.get("status") != "error" else 2


if __name__ == "__main__":
    sys.exit(main())
